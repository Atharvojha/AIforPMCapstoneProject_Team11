"""Build the customer-session workbooks.

    PYTHONPATH=<dir with openpyxl> python3 src/build_customer_workbooks.py

Writes customer/05_customer_evidence.xlsx (capture sheets + facilitator answer key) and
customer/05_recruit_plan_PRIVATE.xlsx (names; never submit). The answer key is derived with the same labelling logic as the
golden set (src/label_items.py), so it is AI-drafted and not human-verified. No formulas: totals come from src/summarise_sessions.py.
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
sys.path.insert(0, HERE)
import label_items as LI  # noqa: E402

F, NAVY = "Arial", "1F3A5F"
HDR = PatternFill("solid", fgColor=NAVY)
INPUT = PatternFill("solid", fgColor="FFF9DB")
GREY = PatternFill("solid", fgColor="EEF2F7")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
PAIR_LAPTOPS = ["L01", "L03", "L04", "L06"]  # Pair A = L01 vs L03 (budget), Pair B = L04 vs L06 (premium)


def sheet(wb, name, headers, widths, rows=(), input_rows=0, example=None, dv=None):
    ws = wb.create_sheet(name)
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill, cell.border = HDR, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 42
    r = 2
    if example:
        for c, v in enumerate(example, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=F, size=10, italic=True)
            cell.alignment, cell.border, cell.fill = WRAP, BORDER, GREY
        r += 1
    for row in rows:
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=F, size=10)
            cell.alignment, cell.border = WRAP, BORDER
        r += 1
    for rr in range(r, r + input_rows):
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=rr, column=c)
            cell.font = Font(name=F, size=10)
            cell.alignment, cell.border, cell.fill = WRAP, BORDER, INPUT
    ws.freeze_panes = "B2"
    for col_letter, options in (dv or {}).items():
        v = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        v.add(f"{col_letter}2:{col_letter}{r + input_rows}")
        ws.add_data_validation(v)
    return ws


def answer_key():
    d = json.load(open(os.path.join(ROOT, "ui", "data.json"), encoding="utf-8"))
    laps = {l["id"]: l for l in d["laptops"]}
    out = []
    ui_words = {"performance": "Performance", "thermals_noise": "Heat and fan noise", "battery": "Battery", "display": "Screen",
                "keyboard_trackpad": "Keyboard and trackpad", "build_weight": "Build and weight", "ports_connectivity": "Ports and Wi-Fi",
                "webcam_mic_speakers": "Webcam, mic and speakers", "ram_storage_upgradeability": "Memory and storage",
                "software_drivers_os": "Software and OS", "after_sales_support": "After-sales support", "value_price": "Value"}
    pretty = lambda a: ui_words.get(a, a)   # same words the student sees on screen
    for lid in PAIR_LAPTOPS:
        revs = [{"review_id": r["id"], "review_text": r["text"]} for r in laps[lid]["reviews"]]
        for p in d["profiles"]:
            lb = LI.label({"item_id": "K", "profile": {"profile_id": p["id"]}, "reviews": revs})
            fmt = lambda dct: "; ".join(f"{pretty(a)} ({', '.join(ids)})" for a, ids in dct.items()) or "none"
            conf = "; ".join(f"{pretty(c['aspect'])} ({', '.join(c['a'])} vs {', '.join(c['b'])})" for c in lb["conflicts"]) or "none"
            out.append((lid, laps[lid]["name"], p["id"], p["label"], fmt(lb["dealbreakers"]), fmt(lb["watchouts"]), conf,
                        ", ".join(pretty(g) for g in lb["gaps"]) or "none", lb["confidence"]))
    return out


def evidence():
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    lines = [
        ("LapFit customer evidence (Team 11)", None),
        ("Purpose", "Deliverable 5: 5+ real students, quoted verbatims, and what changed because of them. One row per participant in Sessions; one row per quote in Verbatims."),
        ("Yellow cells", "Your team fills these in. Grey italic rows marked EX are format examples with invented values: delete them."),
        ("Privacy rule", "Use IDs only (P01, P02, ...). Never write names, phone numbers or emails in this file. Names live in 05_recruit_plan_PRIVATE.xlsx, which is never submitted."),
        ("Real users only", "Teammates and classmates in AI for PMs do not count. The Sessions tab has a 'counts' column: mark N for anyone who does not qualify."),
        ("Quote, don't summarise", "exact_quote must be the participant's words, copied from your notes or recording. Add a time mark so it can be checked."),
        ("Scoring", "Use the Answer_key tab to decide whether a participant 'named the key dealbreaker'. It is AI-drafted and not human-verified: if a participant raises a real concern the key misses, mark Partly and write why in Verbatims."),
        ("Totals", "Run:  python3 src/summarise_sessions.py customer/05_customer_evidence.xlsx   (counts, averages and targets; there are no formulas in this workbook)."),
        ("Targets (set before the sessions)", "At least 4 of 5 students name the key dealbreaker and the main disagreement with the tool; confidence with the tool is higher than with raw reviews; at least 4 of 5 explain 'dealbreaker' correctly in their own words; every theme backed by 2+ participants or marked single."),
        ("Version tested", "Write the UI version in Sessions.version_tested each time. The first version is 'ui v0 rules-engine view'. When the team changes something because of feedback, log it in Changes and record which participants saw the new version."),
    ]
    for i, (a, b) in enumerate(lines, 1):
        ca, cb = ws.cell(row=i, column=1, value=a), ws.cell(row=i, column=2, value=b)
        ca.font = Font(name=F, bold=True, size=14 if i == 1 else 10, color=NAVY if i == 1 else "000000")
        cb.font = Font(name=F, size=10)
        cb.alignment = ca.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A3"].fill = INPUT

    sess_h = ["participant_id", "counts (Y/N)", "session_date", "facilitator", "note_taker", "segment", "shopping_stage", "profile_card_chosen", "profile_fit",
              "order", "baseline_pair", "tool_pair", "baseline_minutes", "baseline_confidence_1to5", "baseline_leaned_to", "baseline_named_key_dealbreaker",
              "baseline_named_main_disagreement", "tool_minutes", "tool_confidence_1to5", "tool_leaned_to", "tool_named_key_dealbreaker",
              "tool_named_main_disagreement", "tool_opened_quotes", "assisted", "dealbreaker_meaning_correct", "trust_vs_own_reading_1to5",
              "would_try_on_real_shortlist", "version_tested", "consent", "recording", "notes"]
    sess_w = [12, 10, 12, 12, 12, 16, 18, 20, 11, 22, 12, 12, 11, 13, 16, 14, 14, 11, 13, 16, 14, 14, 11, 10, 14, 13, 14, 20, 9, 10, 40]
    yn = ["Y", "N"]
    ex = ["EX1", "Y", "2026-11-02", "Aditi", "Sricharan", "CS / coding", "shortlisted 3, buying in a month", "CS undergrad, year 2", "Good", "baseline A, then tool B", "A", "B", 6, 2,
          "Acer", "N", "NA", 7, 4, "MacBook", "Y", "Y", "Y", "N", "Y", 4, "Y", "ui v0 rules-engine view", "Y", "N", "EXAMPLE with invented values: delete this row"]
    ss = sheet(wb, "Sessions", sess_h, sess_w, example=ex, input_rows=20, dv={
        "B": yn, "F": ["CS / coding", "Design / media", "General study", "Budget-first", "Other"], "I": ["Good", "Partial", "None"],
        "P": ["Y", "N", "Partly", "NA"], "Q": ["Y", "N", "Partly", "NA"], "U": ["Y", "N", "Partly", "NA"], "V": ["Y", "N", "Partly", "NA"],
        "W": yn, "X": yn, "Y": ["Y", "N", "Partly"], "AA": yn, "AC": yn, "AD": yn})
    sheet(wb, "Verbatims", ["verbatim_id", "participant_id", "stage", "time_mark", "exact_quote", "what they were doing", "tag", "severity_1to3", "theme_id"],
          [11, 12, 12, 10, 70, 40, 14, 11, 10],
          example=["EX1", "EX1", "Tool", "04:20", "EXAMPLE (invented, delete): I thought 'dealbreaker' meant the laptop was bad overall.", "Reading the red block on the first laptop", "confusion", 2, "T1"],
          input_rows=80, dv={"C": ["Background", "Baseline", "Tool", "Debrief"], "G": ["confusion", "trust", "value", "missing", "bug", "praise", "behaviour"], "H": ["1", "2", "3"]})
    sheet(wb, "Observations", ["obs_id", "participant_id", "screen_element", "what happened", "what they expected", "assisted (Y/N)", "severity_1to3"],
          [9, 12, 22, 55, 40, 12, 11],
          example=["EX1", "EX1", "Evidence sheet", "EXAMPLE (invented, delete): tried to swipe the sheet down on a laptop and it did not move", "It to close", "N", 1],
          input_rows=60, dv={"C": ["Profile cards", "Laptop picker", "Evidence meter", "Review strip", "Dealbreakers block", "Watch-outs block", "Strengths block", "Disagree block", "Evidence gap block", "Claim row", "Evidence sheet", "Source toggle", "Other"], "F": yn, "G": ["1", "2", "3"]})
    sheet(wb, "Themes", ["theme_id", "theme (one line)", "type", "participants (IDs, comma separated)", "how many", "example verbatim_ids", "decision"],
          [9, 55, 16, 30, 10, 22, 45],
          example=["EX1", "EXAMPLE (invented): 'dealbreaker' read as 'bad overall'", "comprehension", "P01, P03", 2, "V001, V007", "Rename and add a one-line definition on the card"],
          input_rows=25, dv={"C": ["usability", "comprehension", "trust", "content", "missing", "other"]})
    sheet(wb, "Changes", ["change_id", "what we changed", "why (verbatim_ids)", "theme_id", "version_after", "date_changed", "who", "participants who saw it", "what happened next", "kept (Y/N)"],
          [9, 50, 20, 9, 22, 12, 12, 22, 45, 10],
          example=["EX1", "EXAMPLE (invented): reworded 'Dealbreakers for you' to 'Problems that matter to you' and added a definition line", "V001, V007", "T1", "ui v1", "2026-11-04", "Prathyusha", "P04, P05", "P04 and P05 both explained it correctly", "Y"],
          input_rows=15, dv={"J": yn})
    key = answer_key()
    sheet(wb, "Answer_key", ["laptop_id", "laptop", "profile_id", "profile", "dealbreakers (review ids)", "watch-outs, single complaint", "reviewers disagree", "not enough evidence on", "evidence confidence"],
          [9, 34, 9, 34, 44, 40, 36, 36, 14], rows=key)
    ak = wb["Answer_key"]
    ak.cell(row=len(key) + 3, column=1, value="AI-DRAFTED, NOT HUMAN-VERIFIED. How to use: if 'dealbreakers' says 'none', then 'named the key dealbreaker' is NA, and a participant who invents one should be noted in Verbatims. Pair A = Acer Aspire 14 vs MSI Modern 14. Pair B = Samsung Galaxy Book4 vs Apple MacBook Air M4.").font = Font(name=F, bold=True, size=10, color="A8301A")
    sc = wb.create_sheet("Scorecard")
    sc.column_dimensions["A"].width = 58; sc.column_dimensions["B"].width = 22; sc.column_dimensions["C"].width = 40
    sc.append(["Measure (fill from summarise_sessions.py)", "Result", "Target"])
    for c in sc[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF", size=10); c.fill = HDR
    for m, t in [("Students who count (real, shopping, not on the team)", "5 or more"), ("Named the key dealbreaker: baseline", "n/a, for comparison"), ("Named the key dealbreaker: tool", "4 of 5 or more"),
                 ("Named the main disagreement: tool", "4 of 5 or more"), ("Average confidence: baseline", ""), ("Average confidence: tool", "higher than baseline"),
                 ("Average minutes: baseline", ""), ("Average minutes: tool", ""), ("Explained 'dealbreaker' correctly", "4 of 5 or more"), ("Needed a hint to use the tool", "as few as possible"),
                 ("Opened at least one quote", "most of them"), ("Would try it on their real shortlist", ""), ("Themes backed by 2+ students", "all changes link to one"), ("Changes made and kept", "")]:
        sc.append([m, None, t])
    for row in sc.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F, size=10); c.border = BORDER; c.alignment = WRAP
        row[1].fill = INPUT
    wb.save(os.path.join(ROOT, "customer", "05_customer_evidence.xlsx"))
    return len(key)


def recruit():
    wb = Workbook()
    ws = wb.active
    ws.title = "Recruit_PRIVATE"
    ws["A1"] = "PRIVATE: contains names. Do NOT submit or share this file. Use participant IDs everywhere else."
    ws["A1"].font = Font(name=F, bold=True, color="A8301A", size=11)
    heads = ["#", "first name or initials", "segment guess", "how you reach them", "contacted on", "session booked for", "status", "assigned ID", "screener ok? (Y/N)", "notes"]
    widths = [5, 20, 16, 30, 13, 18, 13, 11, 12, 40]
    for c, (h, w) in enumerate(zip(heads, widths), 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font, cell.fill, cell.border = Font(name=F, bold=True, color="FFFFFF", size=10), HDR, BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ex = ["EX", "EXAMPLE (invented)", "Budget-first", "hostel WhatsApp group, DM", "2026-10-30", "2026-11-02 17:00", "booked", "P01", "Y", "delete this row"]
    for c, v in enumerate(ex, 1):
        cell = ws.cell(row=4, column=c, value=v)
        cell.font, cell.fill, cell.border, cell.alignment = Font(name=F, size=10, italic=True), GREY, BORDER, WRAP
    for r in range(5, 14):
        for c in range(1, len(heads) + 1):
            cell = ws.cell(row=r, column=c, value=(r - 4) if c == 1 else None)
            cell.font, cell.fill, cell.border = Font(name=F, size=10), INPUT, BORDER
    v = DataValidation(type="list", formula1='"idea,invited,confirmed,done,no-show,declined"', allow_blank=True)
    v.add("G4:G13"); ws.add_data_validation(v)
    ws["A15"] = "Aim to invite 8 people to get 5 completed sessions. Include at least one student from each segment: coding, design/media, general study, budget-first. Fill this in this week."
    ws["A15"].font = Font(name=F, size=10)
    wb.save(os.path.join(ROOT, "customer", "05_recruit_plan_PRIVATE.xlsx"))


if __name__ == "__main__":
    os.makedirs(os.path.join(ROOT, "customer"), exist_ok=True)
    n = evidence()
    recruit()
    print(f"built customer/05_customer_evidence.xlsx ({n} answer-key rows) and customer/05_recruit_plan_PRIVATE.xlsx")

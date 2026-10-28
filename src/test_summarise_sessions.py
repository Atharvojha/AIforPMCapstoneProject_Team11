"""Tests for summarise_sessions.py with invented numbers (a temporary workbook). Run: python3 src/test_summarise_sessions.py"""
import os
import sys
import tempfile
import unittest

from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import summarise_sessions as ss

SESSION_COLS = ["participant_id", "counts (Y/N)", "segment", "baseline_minutes", "baseline_confidence_1to5", "baseline_named_key_dealbreaker",
                "tool_minutes", "tool_confidence_1to5", "tool_named_key_dealbreaker", "tool_named_main_disagreement", "tool_opened_quotes", "assisted",
                "dealbreaker_meaning_correct", "trust_vs_own_reading_1to5", "would_try_on_real_shortlist"]


def make(path, sessions, themes=(), changes=()):
    wb = Workbook()
    ws = wb.active; ws.title = "Sessions"; ws.append(SESSION_COLS)
    for s in sessions:
        ws.append([s.get(c) for c in SESSION_COLS])
    t = wb.create_sheet("Themes"); t.append(["theme_id", "theme (one line)", "type", "participants (IDs, comma separated)"])
    for r in themes: t.append(r)
    c = wb.create_sheet("Changes"); c.append(["change_id", "what we changed", "why (verbatim_ids)", "kept (Y/N)"])
    for r in changes: c.append(r)
    wb.save(path)


def person(i, **kw):
    base = {"participant_id": f"P{i:02d}", "counts (Y/N)": "Y", "segment": "CS / coding", "baseline_minutes": 6, "baseline_confidence_1to5": 2,
            "baseline_named_key_dealbreaker": "N", "tool_minutes": 8, "tool_confidence_1to5": 4, "tool_named_key_dealbreaker": "Y",
            "tool_named_main_disagreement": "Y", "tool_opened_quotes": "Y", "assisted": "N", "dealbreaker_meaning_correct": "Y",
            "trust_vs_own_reading_1to5": 4, "would_try_on_real_shortlist": "Y"}
    base.update(kw); return base


class T(unittest.TestCase):
    def run_it(self, sessions, **kw):
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "e.xlsx"); make(p, sessions, **kw)
            return ss.summarise(p)

    def test_five_good_sessions_meet_the_targets(self):
        o = self.run_it([person(i) for i in range(1, 6)])
        self.assertEqual(o["n"], 5)
        self.assertEqual(o["tool_dealbreaker"], (5, 5)); self.assertEqual(o["baseline_dealbreaker"], (0, 5))
        self.assertAlmostEqual(o["conf_baseline"], 2.0); self.assertAlmostEqual(o["conf_tool"], 4.0)
        self.assertTrue(all(v == "MET" for k, v in o["targets"].items() if "change" not in k), o["targets"])

    def test_examples_non_counting_and_na_are_left_out(self):
        s = [person(1), person(2), {**person(3), "participant_id": "EX1"}, person(4, **{"counts (Y/N)": "N"}), person(5, tool_named_key_dealbreaker="NA"), person(6), person(7)]
        o = self.run_it(s)
        self.assertEqual(o["n"], 5)                     # P01, P02, P05, P06, P07
        self.assertEqual(o["tool_dealbreaker"], (4, 4))  # P05's NA is not in the denominator

    def test_fewer_than_five_is_not_yet_and_misses_are_reported(self):
        self.assertEqual(self.run_it([person(1), person(2)])["targets"]["tool: named the key dealbreaker (4 of 5)"], "NOT YET")
        o = self.run_it([person(i, tool_named_key_dealbreaker="N" if i > 2 else "Y") for i in range(1, 6)])
        self.assertEqual(o["targets"]["tool: named the key dealbreaker (4 of 5)"], "NOT MET")

    def test_confidence_must_actually_rise(self):
        o = self.run_it([person(i, tool_confidence_1to5=2) for i in range(1, 6)])
        self.assertEqual(o["targets"]["confidence with tool above printed reviews"], "NOT MET")

    def test_themes_and_changes_need_evidence(self):
        o = self.run_it([person(1)], themes=[["T1", "x", "usability", "P01, P02"], ["T2", "y", "trust", "P03"], ["EX1", "z", "x", "P01, P02, P03"]],
                        changes=[["C1", "reworded", "V001, V002", "Y"], ["C2", "moved", None, "N"], ["EX1", "e", "V1", "Y"]])
        self.assertEqual((o["themes"], o["themes_2plus"]), (2, 1))
        self.assertEqual((o["changes"], o["changes_kept"], o["changes_without_evidence"]), (2, 1, 1))
        self.assertEqual(o["targets"]["every change links to evidence"], "NOT MET")

    def test_the_real_template_loads_and_is_empty(self):
        root = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "customer", "05_customer_evidence.xlsx")
        if os.path.exists(root):
            self.assertEqual(ss.summarise(root)["n"], 0)  # only the EX1 example row exists, and it must not count


if __name__ == "__main__":
    unittest.main(verbosity=1)

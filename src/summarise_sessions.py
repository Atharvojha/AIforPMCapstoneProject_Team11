"""Totals and target checks for the customer sessions.

    python3 src/summarise_sessions.py customer/05_customer_evidence.xlsx

Reads Sessions, Themes and Changes. Only rows whose participant_id looks like P01, P02... and whose 'counts' column is Y are
included (example rows start with EX and never count). NA answers are left out of the denominators.
"""
import re
import sys

from openpyxl import load_workbook

ID = re.compile(r"^P\d+$")


def rows(ws):
    head = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(v is not None for v in r):
            yield dict(zip(head, r))


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def mean(xs):
    xs = [x for x in xs if x is not None]
    return sum(xs) / len(xs) if xs else None


def rate(vals, good=("Y",)):
    """(successes, denominator) ignoring blanks and NA."""
    vals = [v for v in vals if v not in (None, "", "NA")]
    return sum(v in good for v in vals), len(vals)


def summarise(path):
    wb = load_workbook(path, data_only=True)
    S = [r for r in rows(wb["Sessions"]) if ID.match(str(r.get("participant_id") or "")) and r.get("counts (Y/N)") == "Y"]
    out = {"n": len(S)}
    col = lambda k: [r.get(k) for r in S]
    out["segments"] = {s: col("segment").count(s) for s in set(col("segment")) if s}
    out["baseline_dealbreaker"] = rate(col("baseline_named_key_dealbreaker"))
    out["tool_dealbreaker"] = rate(col("tool_named_key_dealbreaker"))
    out["tool_disagreement"] = rate(col("tool_named_main_disagreement"))
    out["meaning_correct"] = rate(col("dealbreaker_meaning_correct"))
    out["conf_baseline"] = mean([num(x) for x in col("baseline_confidence_1to5")])
    out["conf_tool"] = mean([num(x) for x in col("tool_confidence_1to5")])
    out["min_baseline"] = mean([num(x) for x in col("baseline_minutes")])
    out["min_tool"] = mean([num(x) for x in col("tool_minutes")])
    out["trust"] = mean([num(x) for x in col("trust_vs_own_reading_1to5")])
    out["assisted"] = rate(col("assisted"))
    out["opened_quotes"] = rate(col("tool_opened_quotes"))
    out["would_try"] = rate(col("would_try_on_real_shortlist"))
    T = [r for r in rows(wb["Themes"]) if r.get("theme_id") and not str(r["theme_id"]).startswith("EX")]
    out["themes"] = len(T)
    out["themes_2plus"] = sum(len([x for x in str(r.get("participants (IDs, comma separated)") or "").split(",") if x.strip()]) >= 2 for r in T)
    C = [r for r in rows(wb["Changes"]) if r.get("change_id") and not str(r["change_id"]).startswith("EX")]
    out["changes"] = len(C)
    out["changes_kept"] = sum(r.get("kept (Y/N)") == "Y" for r in C)
    out["changes_without_evidence"] = sum(not r.get("why (verbatim_ids)") for r in C)
    ok = lambda t, need=0.8: "NOT YET" if t[1] < 5 else ("MET" if t[0] / t[1] >= need else "NOT MET")
    cb, ct = out["conf_baseline"], out["conf_tool"]
    out["targets"] = {
        "5 or more real students": "MET" if out["n"] >= 5 else f"NOT YET ({out['n']} so far)",
        "tool: named the key dealbreaker (4 of 5)": ok(out["tool_dealbreaker"]),
        "tool: named the main disagreement (4 of 5)": ok(out["tool_disagreement"]),
        "explained 'dealbreaker' correctly (4 of 5)": ok(out["meaning_correct"]),
        "confidence with tool above printed reviews": "NOT YET" if cb is None or ct is None or out["n"] < 5 else ("MET" if ct > cb else "NOT MET"),
        "every change links to evidence": "NOT YET" if not out["changes"] else ("MET" if out["changes_without_evidence"] == 0 else "NOT MET"),
    }
    return out


def fmt(t):
    return f"{t[0]} of {t[1]}" if t[1] else "no data"


def main():
    o = summarise(sys.argv[1])
    if not o["n"]:
        print("No counting participants yet (rows need an ID like P01 and counts = Y).")
        return
    f = lambda x: "no data" if x is None else f"{x:.1f}"
    print(f"Students who count: {o['n']}   segments: {o['segments']}")
    print(f"Named the key dealbreaker   printed reviews: {fmt(o['baseline_dealbreaker'])}   tool: {fmt(o['tool_dealbreaker'])}")
    print(f"Named the main disagreement tool: {fmt(o['tool_disagreement'])}   explained 'dealbreaker' correctly: {fmt(o['meaning_correct'])}")
    print(f"Confidence 1 to 5   printed: {f(o['conf_baseline'])}   tool: {f(o['conf_tool'])}   | minutes   printed: {f(o['min_baseline'])}   tool: {f(o['min_tool'])}   | trust: {f(o['trust'])}")
    print(f"Needed a hint: {fmt(o['assisted'])}   opened a quote: {fmt(o['opened_quotes'])}   would try on real shortlist: {fmt(o['would_try'])}")
    print(f"Themes: {o['themes']} ({o['themes_2plus']} backed by 2+ students)   changes: {o['changes']} ({o['changes_kept']} kept, {o['changes_without_evidence']} without evidence)")
    print("TARGETS")
    for k, v in o["targets"].items():
        print(f"  {k:48s} {v}")


if __name__ == "__main__":
    main()

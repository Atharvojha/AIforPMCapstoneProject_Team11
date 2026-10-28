"""Validate the golden dataset workbook against the labelling rules, and export it.

Usage:
    python3 src/validate_golden.py 02_golden_dataset_v1.xlsx
    python3 src/validate_golden.py 02_golden_dataset_v1.xlsx --export golden_v1.jsonl

Only items that have a laptop_id are checked (unstarted blueprint rows are counted, not failed).
Exit code 1 if any error is found. Warnings do not fail the run.
"""
import argparse
import json
import re
import sys

from openpyxl import load_workbook

ASPECTS = {
    "performance", "thermals_noise", "battery", "display", "keyboard_trackpad",
    "build_weight", "ports_connectivity", "webcam_mic_speakers",
    "ram_storage_upgradeability", "software_drivers_os", "after_sales_support",
    "value_price", "other",
}
CONFIDENCE = {"High", "Medium", "Low", "Insufficient"}
RID = re.compile(r"R\d{3}")


def rows(ws):
    head = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(v is not None for v in r):
            yield dict(zip(head, r))


def ids(text):
    return set(RID.findall(text or ""))


def lines(text):
    return [ln.strip() for ln in (text or "").splitlines() if ln.strip()]


def check_aspect_lines(label, text, applicable, errors):
    seen = {}
    for ln in lines(text):
        m = re.match(r"^([a-z_]+)\s*:\s*(.+)$", ln)
        if not m or m.group(1) not in ASPECTS:
            errors.append(f"{label}: bad line '{ln}' (use 'aspect: R001, R002' with a valid aspect)")
            continue
        got = ids(m.group(2))
        if not got:
            errors.append(f"{label}: no review ids in '{ln}'")
        if got - applicable:
            errors.append(f"{label}: {sorted(got - applicable)} not in applicable_review_ids")
        seen[m.group(1)] = got
    return seen


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--export", help="write labelled items as JSONL")
    ap.add_argument("--all", action="store_true", help="export all labelled items, not only Frozen")
    a = ap.parse_args()

    wb = load_workbook(a.workbook, data_only=True)
    profiles = {r["profile_id"]: r for r in rows(wb["Profiles"])}
    laptops = {r["laptop_id"]: r for r in rows(wb["Laptops"]) if not str(r["laptop_id"]).startswith("EX")}
    reviews = {r["review_id"]: r for r in rows(wb["Review_Bank"])
               if not str(r["review_id"]).startswith("EX") and r.get("review_text")}
    items = list(rows(wb["Items"]))

    errors, warns, started, assigned, export = [], [], 0, 0, []
    counts = {}
    for it in items:
        iid = it["item_id"]
        if not it.get("laptop_id"):
            continue
        started += 1
        errs = []
        lap = it["laptop_id"]
        if it.get("status") in ("Blueprint", "Assigned", "Needs synthetic"):
            # assigned but not yet labelled: only check that the references resolve
            assigned += 1
            if lap not in laptops or not laptops[lap].get("brand_model"):
                errs.append(f"laptop_id {lap} missing or empty in Laptops tab")
            if it["profile_id"] not in profiles:
                errs.append(f"unknown profile {it['profile_id']}")
            for r_ in sorted(ids(it.get("review_ids"))):
                if r_ not in reviews:
                    errs.append(f"{r_} not in Review_Bank (or has no text)")
                elif reviews[r_]["laptop_id"] != lap:
                    errs.append(f"{r_} belongs to {reviews[r_]['laptop_id']}, not {lap}")
            errors += [f"{iid}: {e}" for e in errs]
            continue
        if lap not in laptops or not laptops[lap].get("brand_model"):
            errs.append(f"laptop_id {lap} missing or empty in Laptops tab")
        if it["profile_id"] not in profiles:
            errs.append(f"unknown profile {it['profile_id']}")
        rv = ids(it.get("review_ids"))
        app_ids, na_ids = ids(it.get("applicable_review_ids")), ids(it.get("not_applicable_review_ids"))
        for r_ in sorted(rv):
            if r_ not in reviews:
                errs.append(f"{r_} not in Review_Bank (or has no text)")
            elif reviews[r_]["laptop_id"] != lap:
                errs.append(f"{r_} belongs to {reviews[r_]['laptop_id']}, not {lap}")
        if app_ids | na_ids != rv or app_ids & na_ids:
            errs.append("applicable + not_applicable must partition review_ids exactly")
        db = check_aspect_lines(f"{iid} dealbreakers", it.get("dealbreakers"), app_ids, errs)
        wo = check_aspect_lines(f"{iid} watchouts", it.get("watchouts_unverified"), app_ids, errs)
        for asp in set(db) & set(wo):
            errs.append(f"aspect {asp} is both a dealbreaker and a watch-out")
        for asp, got in db.items():
            if len(got) < 2:
                warns.append(f"{iid}: dealbreaker '{asp}' has 1 review; allowed only for a verifiable listing/hardware fact (rule b)")
        for asp, got in wo.items():
            if len(got) > 1:
                warns.append(f"{iid}: watch-out '{asp}' has {len(got)} reviews; should it be a dealbreaker?")
        for ln in lines(it.get("conflicts")):
            m = re.match(r"^([a-z_]+)\s*:\s*(.+?)\s+vs\s+(.+?)\s*\|\s*(explained\s*:\s*\w+|unresolved)\s*$", ln)
            if not m or m.group(1) not in ASPECTS:
                errs.append(f"conflict: bad line '{ln}'")
            elif not (ids(m.group(2)) and ids(m.group(3))) or (ids(m.group(2)) | ids(m.group(3))) - app_ids:
                errs.append(f"conflict: both sides need applicable review ids in '{ln}'")
        for g in [x.strip() for x in (it.get("evidence_gaps") or "").split(",") if x.strip()]:
            if g not in ASPECTS:
                errs.append(f"evidence_gaps: unknown aspect '{g}'")
        conf = it.get("expected_confidence")
        if conf not in CONFIDENCE:
            errs.append("expected_confidence missing or invalid")
        if len(rv) < 3 and conf != "Insufficient":
            errs.append("fewer than 3 reviews in total, so confidence must be Insufficient")
        if not it.get("why_correct"):
            errs.append("why_correct is empty")
        if not it.get("labelled_by"):
            errs.append("labelled_by is empty")
        if it.get("status") in ("Second labelled", "Reconciled", "Frozen") and not it.get("second_labeller"):
            errs.append("status needs second_labeller")
        if it.get("status") == "Frozen" and it.get("labels_agree") != "Y":
            errs.append("Frozen items must have labels_agree = Y (reconcile first)")
        errors += [f"{iid}: {e}" for e in errs]
        counts[it["type"]] = counts.get(it["type"], 0) + 1
        if not errs and (a.all or it.get("status") == "Frozen"):
            export.append({
                "item_id": iid, "type": it["type"], "subtype": it["subtype"],
                "profile": profiles[it["profile_id"]], "laptop": laptops[lap],
                "reviews": [{**{k: reviews[r_][k] for k in ("review_id", "stars", "reviewer_config_stated",
                                                             "ownership_duration_stated", "language", "review_text",
                                                             "helpful_votes", "review_date")},
                             "synthetic": str(reviews[r_].get("provenance_note") or "").startswith("SYNTHETIC")}
                            for r_ in sorted(rv)],
                "expected": {
                    "applicable_review_ids": sorted(app_ids), "not_applicable_review_ids": sorted(na_ids),
                    "dealbreakers": {k: sorted(v) for k, v in db.items()},
                    "watchouts": {k: sorted(v) for k, v in wo.items()},
                    "conflicts": lines(it.get("conflicts")),
                    "evidence_gaps": [x.strip() for x in (it.get("evidence_gaps") or "").split(",") if x.strip()],
                    "confidence": conf, "why": it["why_correct"],
                },
            })

    print(f"Items in blueprint: {len(items)} | with laptop: {started} (assigned, unlabelled: {assigned}) | labelled by type: {counts}")
    print(f"Reviews with text: {len(reviews)} | laptops filled: {sum(1 for v in laptops.values() if v.get('brand_model'))}")
    for w in warns:
        print("WARN ", w)
    for e in errors:
        print("ERROR", e)
    print(f"{len(errors)} error(s), {len(warns)} warning(s)")
    labelled = started - assigned
    if labelled < 30:
        print(f"Freeze check: {labelled}/30 items labelled (need 30+)")
    if a.export:
        with open(a.export, "w") as f:
            for e in export:
                f.write(json.dumps(e, ensure_ascii=False, default=str) + "\n")
        print(f"Exported {len(export)} item(s) to {a.export}")
    sys.exit(1 if errors else 0)


if __name__ == "__main__":
    main()

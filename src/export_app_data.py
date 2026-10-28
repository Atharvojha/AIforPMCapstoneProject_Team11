"""Build the files the UI and the eval runners read from the golden workbook.

    PYTHONPATH=<dir with openpyxl> python3 src/export_app_data.py

Writes (relative to the project folder):
  ui/data.json                    profiles, laptops, real reviews and rules-baseline analyses for every profile x laptop
  data/items_inputs.jsonl         the 40 assigned items (profile, laptop, reviews), no labels: input for the product
  data/baseline_outputs.jsonl     rules-baseline output for every item
Synthetic reviews (provenance_note starts with SYNTHETIC) are never shown in the UI.
"""
import json
import os
import re
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
sys.path.insert(0, HERE)
from baseline_rules import analyze  # noqa: E402

WB = os.path.join(ROOT, "02_golden_dataset_v1.xlsx")


def rows(ws):
    head = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(v is not None for v in r):
            yield dict(zip(head, r))


def clean(d):
    return {k: (v if not hasattr(v, "isoformat") else v.isoformat()) for k, v in d.items()}


def main():
    wb = load_workbook(WB, data_only=True)
    profiles = {r["profile_id"]: clean(r) for r in rows(wb["Profiles"])}
    laptops = {r["laptop_id"]: clean(r) for r in rows(wb["Laptops"]) if r["brand_model"] and not str(r["laptop_id"]).startswith("EX")}
    reviews, synthetic = {}, set()
    for r in rows(wb["Review_Bank"]):
        if str(r["review_id"]).startswith("EX") or not r.get("review_text"):
            continue
        reviews[r["review_id"]] = clean(r)
        if str(r.get("provenance_note") or "").startswith("SYNTHETIC"):
            synthetic.add(r["review_id"])

    def rev_list(lid, only_real=True):
        return [reviews[i] for i in sorted(reviews) if reviews[i]["laptop_id"] == lid and not (only_real and i in synthetic)]

    # items for the evals
    items = []
    for it in rows(wb["Items"]):
        if not it.get("laptop_id"):
            continue
        ids = re.findall(r"R\d{3}", it.get("review_ids") or "")
        items.append({"item_id": it["item_id"], "type": it["type"], "subtype": it["subtype"],
                      "profile": profiles[it["profile_id"]], "laptop": laptops[it["laptop_id"]],
                      "reviews": [reviews[i] for i in ids], "status": it.get("status")})
    os.makedirs(os.path.join(ROOT, "data"), exist_ok=True)
    with open(os.path.join(ROOT, "data", "items_inputs.jsonl"), "w", encoding="utf-8") as f:
        for it in items:
            f.write(json.dumps(it, ensure_ascii=False) + "\n")
    with open(os.path.join(ROOT, "data", "baseline_outputs.jsonl"), "w", encoding="utf-8") as f:
        for it in items:
            f.write(json.dumps(analyze(it["profile"], it["laptop"], it["reviews"], it["item_id"]), ensure_ascii=False) + "\n")

    # UI data
    def split_prio(s):
        return [p.strip() for p in re.split(r"\d\)", s or "") if p.strip()]

    ui = {
        "meta": {"reviews_note": "Reviews were hand-picked by our team from Amazon.in and Flipkart pages and have not been independently verified.",
                 "version": "v0 prototype"},
        "profiles": [{
            "id": p["profile_id"], "label": p["label"], "workloads": p["main_workloads"], "mobility": p["mobility"],
            "budget": p["budget_inr"], "priorities": split_prio(p["priorities (ranked)"]),
            "dealbreakers": [d.strip() for d in (p["stated_dealbreakers"] or "").split(";") if d.strip()],
            "software": p["software_needs"], "hours": p["usage_hours"]} for p in profiles.values()],
        "laptops": [], "analyses": {},
    }
    for lid, l in laptops.items():
        revs = rev_list(lid)
        ui["laptops"].append({
            "id": lid, "name": l["brand_model"], "price": l["price_inr_listing"], "tier": l["price_tier"],
            "cpu_gpu": l["cpu_gpu"], "ram_gb": l["ram_gb"], "storage": l["storage"], "display": l["display"],
            "weight_kg": l["weight_kg"], "note": l["known_controversy"],
            "reviews": [{"id": r["review_id"], "stars": r["stars"], "date": r["review_date"], "config": r["reviewer_config_stated"],
                         "lang": r["language"], "text": r["review_text"], "helpful": r["helpful_votes"], "verified": r["verified_purchase"]}
                        for r in revs]})
        for pid, p in profiles.items():
            ui["analyses"][f"{pid}|{lid}"] = {"rules": analyze(p, l, revs, None), "ai": None}
    os.makedirs(os.path.join(ROOT, "ui"), exist_ok=True)
    with open(os.path.join(ROOT, "ui", "data.json"), "w", encoding="utf-8") as f:
        json.dump(ui, f, ensure_ascii=False)
    print(f"{len(items)} items, {len(profiles)} profiles x {len(laptops)} laptops = {len(ui['analyses'])} analyses; "
          f"{len(synthetic)} synthetic reviews excluded from UI")


if __name__ == "__main__":
    main()
